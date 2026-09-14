from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT_PATH = "Wearlytics_Title_Slide.pptx"

WHITE = RGBColor(255, 255, 255)
DEEP_LAVENDER = RGBColor(47, 29, 88)
MEDIUM_LAVENDER = RGBColor(103, 73, 163)
LIGHT_LAVENDER = RGBColor(210, 197, 239)


def add_textbox(slide, left, top, width, height, text, size, bold=False,
                color=WHITE, font="Aptos", align=PP_ALIGN.LEFT,
                margin=0):
    shape = slide.shapes.add_textbox(left, top, width, height)
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.margin_left = margin
    text_frame.margin_right = margin
    text_frame.margin_top = margin
    text_frame.margin_bottom = margin
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.word_wrap = True

    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_rule(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_panel(slide, left, top, width, height):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = MEDIUM_LAVENDER
    shape.line.color.rgb = LIGHT_LAVENDER
    shape.line.width = Pt(1.2)
    return shape


def build_title_slide():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = DEEP_LAVENDER

    add_rule(slide, Inches(0), Inches(0), Inches(0.12), Inches(7.5), MEDIUM_LAVENDER)
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(11.2), Inches(-0.9), Inches(3.0), Inches(3.0)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = MEDIUM_LAVENDER
    circle.line.fill.background()
    circle.fill.transparency = 28

    add_textbox(
        slide, Inches(0.92), Inches(0.62), Inches(7.5), Inches(0.3),
        "S3 MCA MINI PROJECT — SECOND INTERIM PRESENTATION",
        12, bold=True, color=LIGHT_LAVENDER,
    )
    add_rule(slide, Inches(0.92), Inches(1.13), Inches(1.05), Inches(0.06), MEDIUM_LAVENDER)

    add_textbox(
        slide, Inches(0.92), Inches(1.55), Inches(8.9), Inches(0.7),
        "WEARLYTICS", 36, bold=True, color=WHITE,
    )
    add_textbox(
        slide, Inches(0.92), Inches(2.35), Inches(9.8), Inches(0.98),
        "An AI-Powered Smart Wardrobe and\nPersonalized Outfit Recommendation System",
        22, bold=False, color=LIGHT_LAVENDER,
    )

    add_panel(slide, Inches(0.92), Inches(3.75), Inches(9.65), Inches(2.25))

    add_textbox(
        slide, Inches(1.22), Inches(4.05), Inches(5.6), Inches(0.38),
        "ANCHANA R P", 18, bold=True, color=WHITE,
    )
    add_textbox(
        slide, Inches(1.22), Inches(4.48), Inches(5.0), Inches(0.34),
        "Roll No: 23  |  S3 MCA", 15, color=LIGHT_LAVENDER,
    )
    add_textbox(
        slide, Inches(1.22), Inches(4.96), Inches(8.7), Inches(0.58),
        "LOURDES MATHA COLLEGE OF\nSCIENCE AND TECHNOLOGY",
        14, bold=True, color=WHITE,
    )
    add_textbox(
        slide, Inches(6.15), Inches(4.25), Inches(4.0), Inches(0.8),
        "Guide:\nMS. BISMI K. CHARLEYS",
        15, bold=False, color=LIGHT_LAVENDER,
    )

    presentation.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_title_slide()